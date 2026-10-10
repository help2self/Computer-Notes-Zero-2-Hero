# docx -> JSON converter (Computer Notes site). Usage: python build.py <folder-with-docx>
import docx,json,re,sys,glob,html
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph
from docx.text.run import Run
SRC=sys.argv[1] if len(sys.argv)>1 else '.'
DROP=re.compile(r'BRAHMASTRA|Knowledge Notebook|Learning Notes Form|Series\s*·',re.I)
CHS=[('Ch01','💻','Computer Fundamentals','Computer का परिचय, Hardware, Software, भाषाएँ व इतिहास'),
('Ch02','🪟','Windows','Operating System व Windows का उपयोग'),
('Ch03','📝','MS Word','Document बनाना, Formatting, Table, Mail Merge आदि'),
('Ch04','📊','MS Excel','Formula, Function, Chart, Pivot Table आदि'),
('Ch05','🎞️','MS PowerPoint','Presentation बनाना (Basic Level)'),
('Ch06','🌐','Network & Internet','Network, Protocol, Internet, Security')]
E=html.escape
def runs(p): return [Run(r,p) for r in p._p.iter(qn('w:r'))]
def rich(p):
    o=[]
    for r in runs(p):
        t=r.text
        if not t: continue
        s=E(t).replace('\n','<br>'); fn=(r.font.name or '').lower()
        if any(k in fn for k in('consol','courier','mono')): s='<code>'+s+'</code>'
        if r.bold: s='<b>'+s+'</b>'
        if r.italic: s='<i>'+s+'</i>'
        o.append(s)
    return ''.join(o).strip()
def meta(p):
    for r in runs(p):
        if r.text.strip():
            c=None
            try: c=str(r.font.color.rgb) if r.font.color and r.font.color.type else None
            except Exception: pass
            return (r.font.size.pt if r.font.size else 0),bool(r.bold),c
    return 0,False,None
def islist(p): return p._p.pPr is not None and p._p.pPr.numPr is not None
def role(p,txt):
    sz,b,c=meta(p); sn=p.style.name if p.style is not None else ''
    if sn=='Heading 1': return 'item'
    if b and sz>=18: return 'sec'
    if b and 12.5<=sz<=14.5: return 'item'
    if b and 10.9<=sz<=12.5: return 'label'
    if b and 10<=sz<10.9 and c not in(None,'222222') and ord(txt[0])>0x2000 and not '\u0900'<=txt[0]<='\u097f': return 'label'
    return 'p'
def clean(t): return re.sub(r'\s*[—-]\s*Learning Notes.*$','',t).strip()
def node(p):
    h=rich(p)
    if not h: return None
    if islist(p): return {'k':'li','h':h}
    if h.startswith('• '): return {'k':'li','h':h[2:]}
    return {'k':'p','h':h}
def cell(c):
    return '<br>'.join(('• ' if islist(p) else '')+rich(p) for p in c.paragraphs if rich(p))
def tbl(t):
    return {'k':'tb','hd':len(t.rows)>1,'r':[[cell(c) for c in row.cells] for row in t.rows]}
def conv(path,cid,icon,title,desc):
    d=docx.Document(path); secs=[]; chintro=[]; sec=None; item=None; blk=None
    def tgt():
        nonlocal sec,blk
        if sec is None: sec={'t':'','intro':[],'items':[]}; secs.append(sec)
        return item['b'] if item else sec['intro']
    def add(n):
        nonlocal blk
        L=tgt()
        if blk is None or L is not blk[1]: blk=({'l':'','c':'','n':[]},L); L.append(blk[0])
        blk[0]['n'].append(n)
    for el in d.element.body.iterchildren():
        if el.tag==qn('w:p'):
            p=Paragraph(el,d); txt=p.text.strip()
            if not txt or DROP.search(txt): continue
            r=role(p,txt)
            if r=='sec':
                sec={'t':clean(txt),'intro':[],'items':[]}; secs.append(sec); item=None; blk=None
            elif r=='item':
                if sec is None: tgt()
                item={'t':clean(txt),'b':[]}; sec['items'].append(item); blk=None
            elif r=='label':
                L=tgt(); b={'l':E(txt),'c':meta(p)[2] or '','n':[]}; L.append(b); blk=(b,L)
            else:
                n=node(p)
                if n: add(n)
        elif el.tag==qn('w:tbl'):
            t=Table(el,d); L=tgt(); blk=None
            if len(t.rows)==1 and len(t.columns)==1:
                ps=[q for q in t.rows[0].cells[0].paragraphs if q.text.strip()]
                if not ps: continue
                lab=ps[0] if meta(ps[0])[1] else None
                b={'l':E(lab.text.strip()) if lab else '','c':(meta(lab)[2] if lab else '') or '','n':[]}
                for q in (ps[1:] if lab else ps):
                    n=node(q)
                    if n: b['n'].append(n)
                L.append(b)
            else: L.append({'l':'','c':'','n':[tbl(t)]})
    if secs and not secs[0]['items']: chintro=secs[0]['intro']; secs=secs[1:]
    return {'id':cid,'icon':icon,'title':title,'desc':desc,'intro':chintro,'secs':secs}
def plain(h): return html.unescape(re.sub(r'<[^>]+>','',h.replace('<br>',' '))).strip()
idx=[];miss=0
for pre,icon,title,desc in CHS:
    f=glob.glob(f'{SRC}/{pre}*.docx')[0]; cid=pre.lower(); c=conv(f,cid,icon,title,desc)
    blob=json.dumps(c,ensure_ascii=False); txt=plain(html.unescape(blob.replace('\\"','"')))
    d=docx.Document(f); allp=[q.text for q in d.paragraphs]+[q.text for t in d.tables for r in t.rows for ce in r.cells for q in ce.paragraphs]
    bad=[x for x in set(allp) if x.strip() and not DROP.search(x) and re.sub(r'\s+','',clean(x)) not in re.sub(r'\s+','',txt)]
    miss+=len(bad); [print('  MISSING?',cid,x[:60]) for x in bad[:4]]
    json.dump(c,open(cid+'.json','w'),ensure_ascii=False,separators=(',',':'))
    ni=sum(len(s['items']) for s in c['secs']); idx.append({'id':cid,'icon':icon,'title':title,'desc':desc,'secs':len(c['secs']),'items':ni})
    print(cid,'sections',len(c['secs']),'items',ni,'chars',len(blob))
json.dump(idx,open('index.json','w'),ensure_ascii=False)
print('unmatched paragraphs:',miss)
